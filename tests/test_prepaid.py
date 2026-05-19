from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from build_amortization import build


def test_build_amortization_creates_formula_cells(tmp_path: Path):
    items = pd.DataFrame(
        [
            {
                "item_id": "P001",
                "vendor": "Acme SaaS",
                "invoice_number": "INV-1001",
                "invoice_date": "2026-04-01",
                "amount": 12_000.00,
                "service_start": "2026-04-01",
                "service_end": "2027-03-31",
                "expense_account": "6400",
            }
        ]
    )
    out = tmp_path / "prepaid.xlsx"
    result = build(items, out, period_close="2026-04-30")
    assert result.schedule_rows == 12
    assert out.exists()

    wb = load_workbook(out)
    sch = wb["Schedule"]
    # Non-final months should be formula referencing Inputs
    first_amt = sch.cell(row=2, column=4).value
    assert isinstance(first_amt, str) and first_amt.startswith("="), (
        f"Expected formula in Schedule!D2, got {first_amt!r}"
    )
    assert "Inputs!" in first_amt

    # Final month should reference SUMIFS to absorb rounding
    last_row = sch.max_row
    last_amt = sch.cell(row=last_row, column=4).value
    assert "SUMIFS" in last_amt


def test_build_amortization_rejects_missing_service_period(tmp_path: Path):
    items = pd.DataFrame(
        [
            {
                "item_id": "P002",
                "vendor": "Vendor B",
                "invoice_number": "INV-2002",
                "invoice_date": "2026-04-01",
                "amount": 5_000.00,
                "service_start": None,
                "service_end": None,
                "expense_account": "6400",
            }
        ]
    )
    out = tmp_path / "bad.xlsx"
    try:
        build(items, out)
    except ValueError as exc:
        assert "service period" in str(exc).lower()
    else:
        raise AssertionError("Expected ValueError for missing service period")
