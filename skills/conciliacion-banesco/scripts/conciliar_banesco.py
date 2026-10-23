"""
Conciliación bancaria Banesco: estado de cuenta (transcrito) vs libro de banco.

Pasos:
  1. Carga la transcripción y el resumen (banesco_parse).
  2. Valida la integridad (banesco_validate). Con errores de cuadre NO concilia:
     escribe el libro de trabajo con la pestaña Validacion y termina con código 1.
  3. Clasifica conceptos y empareja comisiones (banesco_classify).
  4. Si hay libro: empareja banco vs libro con bank-rec/scripts/match_transactions
     sobre todo el rango de meses a la vez, y para cada cierre de mes calcula:

        Saldo según estado de cuenta
        + depósitos en tránsito          (en libro, aún no en banco al cierre)
        - cheques / pagos en tránsito    (en libro, aún no en banco al cierre)
        = saldo bancario conciliado
        Saldo según libro
        + notas de crédito no registradas (en banco, aún no en libro al cierre)
        - notas de débito no registradas  (en banco, aún no en libro al cierre)
        = saldo libro ajustado
        Diferencia = conciliado - ajustado  (debe ser 0; nunca se fuerza)

     Una partida emparejada con fechas a ambos lados del cierre (p. ej. depósito
     registrado el 31 y acreditado el 2) es partida en tránsito de ese cierre.
     Las coincidencias probables (MEDIUM) no se confirman: siguen como partidas.
  5. Propone asientos por comisiones bancarias no registradas (JEs_Proposed),
     que deben pasar por je-review antes de contabilizarse.

Libro de banco (CSV/XLSX), columnas aceptadas:
    fecha | date
    descripcion | concepto | description
    referencia | reference                 (opcional)
    monto | amount  (con signo: + entra al banco, - sale)  o bien  debe/haber
    (en el libro de banco el debe aumenta el saldo y el haber lo disminuye)

CLI:
    python conciliar_banesco.py --movimientos movs.csv --resumen resumen.csv \\
        [--libro libro.csv --saldo-inicial-libro 202909,91] \\
        [--cuenta-banco 1010 --cuenta-gasto 6900] [--comisiones bruto|neto] \\
        --out conciliacion.xlsx
"""
from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path
from typing import Iterable

import pandas as pd

HERE = Path(__file__).resolve().parent
BANK_REC_SCRIPTS = HERE.parents[1] / "bank-rec" / "scripts"
if str(BANK_REC_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(BANK_REC_SCRIPTS))

from age_items import age  # noqa: E402  (bank-rec)
from match_transactions import match  # noqa: E402  (bank-rec)

from banesco_classify import (  # noqa: E402
    CATEGORIAS_GASTO_BANCARIO,
    clasificar,
    resumen_mensual,
    tabla_comisiones,
)
from banesco_parse import load_resumen, load_transcripcion, parse_ve_amount  # noqa: E402
from banesco_validate import TOLERANCIA, hay_errores, validar  # noqa: E402

OLD_ITEM_DAYS = 60
NUM_FMT = "#,##0.00;(#,##0.00)"


# ---------------------------------------------------------------------------
# Libro
# ---------------------------------------------------------------------------

_ALIAS = {
    "fecha": "date",
    "descripcion": "description",
    "descripción": "description",
    "concepto": "description",
    "referencia": "reference",
    "ref": "reference",
    "monto": "amount",
}


def load_libro(source: Path | str | pd.DataFrame) -> pd.DataFrame:
    if isinstance(source, pd.DataFrame):
        raw = source.copy()
    else:
        path = Path(source)
        raw = pd.read_csv(path, dtype=str, keep_default_na=False) if path.suffix.lower() == ".csv" \
            else pd.read_excel(path, dtype=str).fillna("")
    raw.columns = [_ALIAS.get(str(c).strip().lower(), str(c).strip().lower()) for c in raw.columns]
    if "date" not in raw.columns or "description" not in raw.columns:
        raise ValueError("libro sin columnas fecha/descripcion")
    if "amount" in raw.columns:
        amount = [parse_ve_amount(v) or 0.0 for v in raw["amount"]]
    elif {"debe", "haber"} <= set(raw.columns):
        amount = [(parse_ve_amount(d) or 0.0) - (parse_ve_amount(h) or 0.0) for d, h in zip(raw["debe"], raw["haber"])]
    else:
        raise ValueError("libro sin monto: use 'monto' con signo o 'debe'/'haber'")
    out = pd.DataFrame(
        {
            "date": pd.to_datetime(raw["date"], dayfirst=True, errors="coerce"),
            "description": raw["description"].astype(str),
            "amount": pd.Series(amount, dtype=float).round(2),
            "reference": raw["reference"].astype(str) if "reference" in raw.columns else "",
        }
    )
    if out["date"].isna().any():
        raise ValueError(f"libro: {int(out['date'].isna().sum())} fila(s) con fecha ilegible")
    return out


# ---------------------------------------------------------------------------
# Banco para el emparejamiento
# ---------------------------------------------------------------------------

def banco_para_match(clas: pd.DataFrame, modo_comisiones: str) -> pd.DataFrame:
    """Una fila por movimiento bancario (bruto) o comisión sumada a su principal (neto)."""
    df = clas.copy()
    df["amount"] = (df["abono"] - df["cargo"]).round(2)
    if modo_comisiones == "neto":
        emparejadas = df.loc[df["es_comision"] & df["principal_linea"].notna()]
        for _, c in emparejadas.iterrows():
            mask = (df["periodo"] == c["periodo"]) & (df["linea"] == c["principal_linea"])
            df.loc[mask, "amount"] = (df.loc[mask, "amount"] - c["monto"]).round(2)
        df = df.drop(emparejadas.index)
    return pd.DataFrame(
        {
            "date": df["fecha"],
            "description": df["concepto"],
            "amount": df["amount"],
            "reference": df["referencia"],
            "periodo": df["periodo"].astype(str),
            "linea": df["linea"],
            "categoria": df["categoria"],
            "id_banco": df["periodo"].astype(str) + "#" + df["linea"].astype(str),
        }
    ).reset_index(drop=True)


# ---------------------------------------------------------------------------
# Partidas en conciliación por cierre
# ---------------------------------------------------------------------------

def _clasificar_banco(categoria: str, amount: float) -> tuple[str, str, str]:
    if categoria in CATEGORIAS_GASTO_BANCARIO:
        return "Nota de débito no registrada (comisión)", "BANK_FEE_NOT_BOOKED", "READY_FOR_REVIEW"
    if categoria == "pago_seniat":
        return "Pago SENIAT no registrado en libro", "UNKNOWN", "MANUAL_REVIEW_REQUIRED"
    if amount > 0:
        return "Nota de crédito no registrada", "UNKNOWN", "MANUAL_REVIEW_REQUIRED"
    return "Nota de débito no registrada", "UNKNOWN", "MANUAL_REVIEW_REQUIRED"


def partidas_por_cierre(
    result, banco: pd.DataFrame, libro: pd.DataFrame, cierre: pd.Timestamp, inicio: pd.Timestamp
) -> pd.DataFrame:
    """Partidas abiertas a la fecha de cierre (lado banco y lado libro)."""
    # las filas emparejadas de bank-rec no traen id_banco: se recupera por fecha, monto, ref y concepto
    por_clave = {
        (pd.Timestamp(r["date"]), round(float(r["amount"]), 2), str(r["reference"]), str(r["description"])): r
        for _, r in banco.iterrows()
    }
    filas: list[dict] = []

    def add(lado, fecha, desc, amount, ref, contraparte_fecha=None):
        id_banco = ""
        if lado == "banco":
            fila = por_clave.get((pd.Timestamp(fecha), round(float(amount), 2), str(ref), str(desc)))
            cat = "otro" if fila is None else fila["categoria"]
            id_banco = "" if fila is None else fila["id_banco"]
            partida, estado, flag = _clasificar_banco(cat, amount)
        else:
            cat = ""
            if amount > 0:
                partida, estado = "Depósito en tránsito", "DEPOSIT_IN_TRANSIT"
            else:
                partida, estado = "Cheque / pago en tránsito", "OUTSTANDING_CHECK"
            flag = "READY_FOR_REVIEW"
        filas.append(
            {
                "cierre": cierre.date().isoformat(),
                "lado": lado,
                "partida": partida,
                "estado_rec": estado,
                "flag": flag,
                "date": pd.Timestamp(fecha),
                "description": desc,
                "amount": round(float(amount), 2),
                "reference": ref,
                "categoria": cat,
                "id_banco": id_banco,
                "se_regulariza": "" if contraparte_fecha is None else pd.Timestamp(contraparte_fecha).date().isoformat(),
            }
        )

    for _, m in result.exact_matches.iterrows():
        bd, kd = pd.Timestamp(m["bank_date"]), pd.Timestamp(m["book_date"])
        if kd <= cierre < bd:
            add("libro", kd, m["book_description"], m["book_amount"], m["book_reference"], bd)
        elif bd <= cierre < kd:
            add("banco", bd, m["bank_description"], m["bank_amount"], m["bank_reference"], kd)
    for _, m in result.probable_matches.iterrows():
        bd, kd = pd.Timestamp(m["bank_date"]), pd.Timestamp(m["book_date"])
        if bd <= cierre:
            add("banco", bd, m["bank_description"], m["bank_amount"], m["bank_reference"])
        if kd <= cierre:
            add("libro", kd, m["book_description"], m["book_amount"], m["book_reference"])
    for _, r in result.unmatched_bank.iterrows():
        if pd.Timestamp(r["date"]) <= cierre:
            add("banco", r["date"], r["description"], r["amount"], r["reference"])
    for _, r in result.unmatched_book.iterrows():
        if pd.Timestamp(r["date"]) <= cierre:
            add("libro", r["date"], r["description"], r["amount"], r["reference"])

    out = pd.DataFrame(filas)
    if out.empty:
        return out
    # age() de bank-rec escribe su propia columna "flag" (OLD_ITEM): se preserva la nuestra
    out = out.rename(columns={"flag": "_flag"})
    out = age(out, period_end=cierre, date_col="date", old_item_days=OLD_ITEM_DAYS)
    out["antigua"] = out["flag"] == "OLD_ITEM"
    out["flag"] = out["_flag"].where(~out["antigua"], "MANUAL_REVIEW_REQUIRED")
    out = out.drop(columns="_flag")
    _marcar_duplicados_libro(out, libro)
    return out


def _marcar_duplicados_libro(partidas: pd.DataFrame, libro: pd.DataFrame) -> None:
    con_ref = libro.loc[libro["reference"].astype(str).str.strip() != ""]
    dup = con_ref.duplicated(subset=["date", "amount", "reference"], keep=False)
    claves = {(pd.Timestamp(r["date"]), r["amount"], str(r["reference"])) for _, r in con_ref.loc[dup].iterrows()}
    for i, r in partidas.loc[partidas["lado"] == "libro"].iterrows():
        if (pd.Timestamp(r["date"]), r["amount"], str(r["reference"])) in claves:
            partidas.at[i, "estado_rec"] = "DUPLICATE_ENTRY"
            partidas.at[i, "flag"] = "MANUAL_REVIEW_REQUIRED"
            partidas.at[i, "partida"] = "Posible registro duplicado en libro"


# ---------------------------------------------------------------------------
# Resumen y asientos
# ---------------------------------------------------------------------------

def saldos_banco(movs: pd.DataFrame, resumen: pd.DataFrame | None) -> dict:
    out = {}
    for periodo, g in movs.groupby("periodo", sort=True):
        g = g.sort_values("linea")
        det = g.loc[~g["es_saldo_anterior"]]
        ant = g.loc[g["es_saldo_anterior"]]
        out[periodo] = {
            "inicial": float(ant.iloc[0]["saldo"]) if not ant.empty else None,
            "final": float(det.iloc[-1]["saldo"]) if not det.empty else None,
        }
    if resumen is not None:
        for _, r in resumen.iterrows():
            d = out.setdefault(r["periodo"], {"inicial": None, "final": None})
            d["inicial"] = d["inicial"] if d["inicial"] is not None else r["saldo_inicial"]
            d["final"] = r["saldo_final"] if r["saldo_final"] is not None else d["final"]
    return out


def tabla_resumen(periodos, saldos, partidas, libro, saldo_inicial_libro, inicio) -> pd.DataFrame:
    filas = []
    for p in periodos:
        cierre = p.end_time.normalize()
        pp = partidas.loc[partidas["cierre"] == cierre.date().isoformat()] if not partidas.empty else partidas
        libro_lado = pp.loc[pp["lado"] == "libro"] if not pp.empty else pp
        banco_lado = pp.loc[pp["lado"] == "banco"] if not pp.empty else pp
        dit = float(libro_lado.loc[libro_lado["amount"] > 0, "amount"].sum()) if not pp.empty else 0.0
        oc = float(-libro_lado.loc[libro_lado["amount"] < 0, "amount"].sum()) if not pp.empty else 0.0
        nc = float(banco_lado.loc[banco_lado["amount"] > 0, "amount"].sum()) if not pp.empty else 0.0
        nd = float(-banco_lado.loc[banco_lado["amount"] < 0, "amount"].sum()) if not pp.empty else 0.0
        saldo_libro = None
        if saldo_inicial_libro is not None and libro is not None:
            dentro = libro.loc[(libro["date"] >= inicio) & (libro["date"] <= cierre), "amount"]
            saldo_libro = round(saldo_inicial_libro + float(dentro.sum()), 2)
        filas.append(
            {
                "periodo": str(p),
                "cierre": cierre.date().isoformat(),
                "saldo_banco": saldos.get(p, {}).get("final"),
                "depositos_transito": round(dit, 2),
                "pagos_transito": round(oc, 2),
                "saldo_libro": saldo_libro,
                "notas_credito_no_reg": round(nc, 2),
                "notas_debito_no_reg": round(nd, 2),
            }
        )
    out = pd.DataFrame(filas)
    for c in ("saldo_banco", "saldo_libro"):
        out[c] = pd.to_numeric(out[c], errors="coerce")
    out["banco_conciliado"] = (out["saldo_banco"] + out["depositos_transito"] - out["pagos_transito"]).round(2)
    out["libro_ajustado"] = (out["saldo_libro"] + out["notas_credito_no_reg"] - out["notas_debito_no_reg"]).round(2)
    out["diferencia"] = (out["banco_conciliado"] - out["libro_ajustado"]).round(2)
    return out


def proponer_asientos(
    comisiones: pd.DataFrame, cuenta_banco: str, cuenta_gasto: str, preparer: str, con_libro: bool
) -> pd.DataFrame:
    """Un asiento por período: Dr gasto bancario / Cr banco por las comisiones dadas."""
    filas = []
    if comisiones.empty:
        return pd.DataFrame(columns=["JE_ID", "Date", "Account", "Debit", "Credit", "Memo", "Support Reference", "Preparer", "Reviewer"])
    n = 0
    for periodo, g in comisiones.groupby("periodo", sort=True):
        n += 1
        total = round(float(g["monto"].sum()), 2)
        if total <= 0:
            continue
        detalle = ", ".join(f"{cat.replace('comision_', '')} {len(x)}" for cat, x in g.groupby("categoria"))
        base = "no registradas en libro" if con_libro else "según estado de cuenta (confirmar que no estén ya registradas en libro)"
        memo = (
            f"Registrar comisiones bancarias Banesco {periodo} {base}: {len(g)} cargos ({detalle}), "
            f"total Bs. {total:,.2f}"
        )
        fecha = pd.Period(periodo, freq="M").end_time.date().isoformat()
        soporte = f"Estado de cuenta Banesco {periodo}; workpaper pestaña Comisiones"
        je_id = f"JE_{n:03d}"
        filas.append(dict(JE_ID=je_id, Date=fecha, Account=cuenta_gasto, Debit=total, Credit=None, Memo=memo,
                          **{"Support Reference": soporte}, Preparer=preparer, Reviewer=""))
        filas.append(dict(JE_ID=je_id, Date=fecha, Account=cuenta_banco, Debit=None, Credit=total, Memo=memo,
                          **{"Support Reference": soporte}, Preparer=preparer, Reviewer=""))
    return pd.DataFrame(filas)


# ---------------------------------------------------------------------------
# Libro de trabajo
# ---------------------------------------------------------------------------

def _evidencia(resumen_rec: pd.DataFrame, movs_path: str, libro_path: str) -> pd.DataFrame:
    filas = []
    for _, r in resumen_rec.iterrows():
        p = r["periodo"]
        filas += [
            (p, "Saldo según estado de cuenta", r["saldo_banco"], f"{movs_path} / resumen {p}", "Último SALDOS del detalle = 'Final' del resumen", ""),
            (p, "Depósitos en tránsito", r["depositos_transito"], "Reconciling_Items", f"SUMA libro, amount>0, cierre {r['cierre']}", ""),
            (p, "Cheques / pagos en tránsito", r["pagos_transito"], "Reconciling_Items", f"SUMA libro, amount<0, cierre {r['cierre']}", ""),
            (p, "Saldo según libro", r["saldo_libro"], libro_path or "no provisto", "saldo inicial libro + movimientos del libro hasta el cierre", ""),
            (p, "Notas de crédito no registradas", r["notas_credito_no_reg"], "Reconciling_Items", f"SUMA banco, amount>0, cierre {r['cierre']}", ""),
            (p, "Notas de débito no registradas", r["notas_debito_no_reg"], "Reconciling_Items", f"SUMA banco, amount<0, cierre {r['cierre']}", ""),
            (p, "Diferencia no conciliada", r["diferencia"], "Summary", "banco conciliado - libro ajustado", ""),
        ]
    return pd.DataFrame(filas, columns=["Periodo", "Concepto", "Valor", "Fuente", "Calculo", "Notas"])


def _escribir_summary(ws, resumen_rec: pd.DataFrame, tolerancia: float) -> None:
    from openpyxl.styles import Font

    etiquetas = [
        ("Saldo según estado de cuenta", "saldo_banco"),
        ("Más: depósitos en tránsito", "depositos_transito"),
        ("Menos: cheques / pagos en tránsito", "pagos_transito"),
        ("Saldo bancario conciliado", None),
        ("Saldo según libro", "saldo_libro"),
        ("Más: notas de crédito no registradas", "notas_credito_no_reg"),
        ("Menos: notas de débito no registradas", "notas_debito_no_reg"),
        ("Saldo libro ajustado", None),
        ("Diferencia no conciliada", None),
        ("Estado", None),
    ]
    ws.cell(row=1, column=1, value="Conciliación bancaria Banesco (Bs.)").font = Font(bold=True)
    ws.cell(row=2, column=1, value="Período")
    for j, (_, r) in enumerate(resumen_rec.iterrows(), start=2):
        col = ws.cell(row=2, column=j).column_letter
        ws.cell(row=2, column=j, value=r["periodo"]).font = Font(bold=True)
        for i, (label, key) in enumerate(etiquetas, start=3):
            ws.cell(row=i, column=1, value=label)
            if key is not None:
                v = r[key]
                ws.cell(row=i, column=j, value=None if pd.isna(v) else float(v)).font = Font(color="0000FF")
        ws.cell(row=6, column=j, value=f"={col}3+{col}4-{col}5")
        if pd.isna(r["saldo_libro"]):
            ws.cell(row=10, column=j, value="n/d")
            ws.cell(row=11, column=j, value="n/d")
            ws.cell(row=12, column=j, value="SIN LIBRO: conciliación incompleta").font = Font(color="FF0000")
        else:
            ws.cell(row=10, column=j, value=f"={col}7+{col}8-{col}9")
            ws.cell(row=11, column=j, value=f"={col}6-{col}10").font = Font(bold=True, underline="double")
            ws.cell(row=12, column=j, value=f'=IF(ABS({col}11)<={tolerancia},"CONCILIADA","NO CONCILIA")')
        for i in range(3, 12):
            ws.cell(row=i, column=j).number_format = NUM_FMT
    ws.column_dimensions["A"].width = 42


def escribir_workbook(out: Path, hojas: dict[str, pd.DataFrame], resumen_rec: pd.DataFrame | None, tolerancia: float) -> None:
    with pd.ExcelWriter(out, engine="openpyxl") as xl:
        for nombre, df in hojas.items():
            if nombre == "Summary":
                continue
            df = df.copy()
            for c in df.columns:
                if isinstance(df[c].dtype, pd.PeriodDtype):
                    df[c] = df[c].astype(str)
            df.to_excel(xl, sheet_name=nombre, index=False)
        if resumen_rec is not None:
            ws = xl.book.create_sheet("Summary", index=1)
            _escribir_summary(ws, resumen_rec, tolerancia)


def conciliar(
    movs_path: Path,
    resumen_path: Path | None,
    out: Path,
    libro_path: Path | None = None,
    saldo_inicial_libro: float | None = None,
    cuenta_banco: str | None = None,
    cuenta_gasto: str = "6900",
    modo_comisiones: str = "bruto",
    preparer: str = "Claude",
    tolerancia: float = TOLERANCIA,
) -> int:
    movs = load_transcripcion(movs_path)
    resumen = load_resumen(resumen_path) if resumen_path else None
    hallazgos = validar(movs, resumen, tolerancia)
    periodos = sorted(movs["periodo"].unique())
    notas: list[dict] = [
        {"tipo": "validacion", "flag": r["flag"], "detalle": f"{r['prueba']} {r['periodo']} línea {r['linea']}: {r['nota']}"}
        for _, r in hallazgos.iterrows()
    ]
    cover = pd.DataFrame(
        {
            "campo": ["Entregable", "Banco", "Moneda", "Períodos", "Libro de banco", "Cuenta contable banco",
                      "Modo comisiones", "Preparado por", "Fecha", "Estado"],
            "valor": ["Conciliación bancaria", "Banesco Banco Universal", "Bs. (VES)",
                      f"{periodos[0]} a {periodos[-1]}", str(libro_path or "NO PROVISTO"), cuenta_banco or "NO INDICADA",
                      modo_comisiones, preparer, date.today().isoformat(), ""],
        }
    )

    if hay_errores(hallazgos):
        cover.loc[cover["campo"] == "Estado", "valor"] = "REJECT: la transcripción no cuadra; corregir antes de conciliar"
        escribir_workbook(out, {"Cover": cover, "Source_Bank": movs, "Validacion": hallazgos,
                                "Review_Notes": pd.DataFrame(notas)}, None, tolerancia)
        print(f"La transcripción tiene errores de cuadre ({int((hallazgos['flag'] == 'MANUAL_REVIEW_REQUIRED').sum())}). Ver {out}, pestaña Validacion.")
        return 1

    clas = clasificar(movs)
    comisiones = tabla_comisiones(clas)
    saldos = saldos_banco(movs, resumen)
    inicio = periodos[0].start_time.normalize()
    hojas: dict[str, pd.DataFrame] = {
        "Cover": cover,
        "Summary": pd.DataFrame(),
        "Source_Bank": movs,
        "Validacion": hallazgos,
        "Clasificacion": clas,
        "Comisiones": comisiones,
        "Resumen_Mensual": resumen_mensual(clas),
    }

    libro = None
    partidas = pd.DataFrame()
    if libro_path is not None:
        libro = load_libro(libro_path)
        banco = banco_para_match(clas, modo_comisiones)
        result = match(banco[["date", "description", "amount", "reference"]], libro)
        partidas = pd.concat(
            [partidas_por_cierre(result, banco, libro, p.end_time.normalize(), inicio) for p in periodos],
            ignore_index=True,
        )
        hojas.update(
            {
                "Source_GL": libro,
                "Matched": result.exact_matches,
                "Probable_Matches": result.probable_matches,
                "Unmatched_Bank": result.unmatched_bank,
                "Unmatched_Book": result.unmatched_book,
                "Reconciling_Items": partidas,
            }
        )
        if saldo_inicial_libro is None:
            notas.append({"tipo": "insumo", "flag": "NEEDS_CLARIFICATION",
                          "detalle": "Falta el saldo inicial del libro al día anterior al primer período: no se calcula el saldo según libros."})
        if len(result.probable_matches):
            notas.append({"tipo": "match", "flag": "NEEDS_CLARIFICATION",
                          "detalle": f"{len(result.probable_matches)} coincidencias probables sin confirmar (pestaña Probable_Matches)."})
        unknown = 0 if partidas.empty else int((partidas["estado_rec"] == "UNKNOWN").sum())
        if len(partidas) and unknown / len(partidas) > 0.30:
            notas.append({"tipo": "calidad", "flag": "MANUAL_REVIEW_REQUIRED",
                          "detalle": f"{unknown} de {len(partidas)} partidas sin clasificar (>30 %): revisar si el libro corresponde a esta cuenta y período."})
        no_reg = partidas.loc[(partidas.get("estado_rec") == "BANK_FEE_NOT_BOOKED")] if not partidas.empty else partidas
        # una comisión no registrada aparece en todos los cierres posteriores: se toma su primer cierre
        if not no_reg.empty:
            no_reg = no_reg.sort_values("cierre").drop_duplicates(subset=["id_banco"])
            je_base = pd.DataFrame({"periodo": no_reg["date"].dt.to_period("M").astype(str),
                                    "categoria": no_reg["categoria"], "monto": -no_reg["amount"]})
        else:
            je_base = pd.DataFrame(columns=["periodo", "categoria", "monto"])
        con_libro = True
    else:
        notas.append({"tipo": "insumo", "flag": "NEEDS_CLARIFICATION",
                      "detalle": "No se recibió libro de banco: solo se validó y clasificó el estado de cuenta. Pedir el mayor de la cuenta para conciliar."})
        je_base = comisiones.assign(periodo=comisiones["periodo"].astype(str))[["periodo", "categoria", "monto"]]
        con_libro = False

    if cuenta_banco:
        hojas["JEs_Proposed"] = proponer_asientos(je_base, cuenta_banco, cuenta_gasto, preparer, con_libro)
    else:
        notas.append({"tipo": "insumo", "flag": "NEEDS_CLARIFICATION",
                      "detalle": "Sin número de cuenta contable del banco (--cuenta-banco): no se proponen asientos."})

    resumen_rec = tabla_resumen(periodos, saldos, partidas, libro, saldo_inicial_libro, inicio)
    for _, r in resumen_rec.iterrows():
        if pd.notna(r["diferencia"]) and abs(r["diferencia"]) > tolerancia:
            notas.append({"tipo": "resultado", "flag": "MANUAL_REVIEW_REQUIRED",
                          "detalle": f"{r['periodo']}: diferencia no conciliada Bs. {r['diferencia']:,.2f}. No se ajustó."})
    for _, r in clas.loc[clas["flag"] != ""].iterrows():
        notas.append({"tipo": "clasificacion", "flag": r["flag"],
                      "detalle": f"{r['periodo']} línea {r['linea']} ref {r['referencia']} {r['concepto']} Bs. {r['monto']:,.2f}: {r['tratamiento']}"})

    estado = "SIN LIBRO: validado y clasificado" if libro is None else (
        "CONCILIADA" if resumen_rec["diferencia"].abs().le(tolerancia).all() else "NO CONCILIA / INCOMPLETA")
    cover.loc[cover["campo"] == "Estado", "valor"] = estado
    hojas["Cover"] = cover
    hojas["Evidence"] = _evidencia(resumen_rec, str(movs_path), str(libro_path or ""))
    hojas["Review_Notes"] = pd.DataFrame(notas, columns=["tipo", "flag", "detalle"])
    escribir_workbook(out, hojas, resumen_rec, tolerancia)

    with pd.option_context("display.width", 200):
        print(resumen_rec.to_string(index=False))
    print(f"Estado: {estado}. Libro de trabajo -> {out}")
    return 0


def _cli(argv: Iterable[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Conciliación bancaria Banesco vs libro de banco.")
    p.add_argument("--movimientos", type=Path, required=True, help="Transcripción CSV/XLSX del estado de cuenta")
    p.add_argument("--resumen", type=Path, help="CSV con el RESUMEN de cada período")
    p.add_argument("--libro", type=Path, help="Libro de banco / mayor de la cuenta")
    p.add_argument("--saldo-inicial-libro", help="Saldo del libro al día anterior al primer período (1.234,56)")
    p.add_argument("--cuenta-banco", help="Cuenta contable del banco (firm-config)")
    p.add_argument("--cuenta-gasto", default="6900", help="Cuenta de gastos bancarios (firm-config)")
    p.add_argument("--comisiones", choices=["bruto", "neto"], default="bruto",
                   help="bruto: el libro registra cobro y comisión por separado; neto: registra el neto")
    p.add_argument("--preparer", default="Claude")
    p.add_argument("--tolerancia", type=float, default=TOLERANCIA)
    p.add_argument("--out", type=Path, required=True)
    a = p.parse_args(argv)
    return conciliar(
        a.movimientos, a.resumen, a.out, a.libro,
        parse_ve_amount(a.saldo_inicial_libro) if a.saldo_inicial_libro else None,
        a.cuenta_banco, a.cuenta_gasto, a.comisiones, a.preparer, a.tolerancia,
    )


if __name__ == "__main__":
    sys.exit(_cli())
